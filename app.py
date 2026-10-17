import os
import csv
import json
import mimetypes
import time
import uuid
import traceback
from pathlib import Path
from datetime import datetime, timezone

import gradio as gr
import requests
from google import genai
from google.genai import types


# =========================================================
# 1. CONFIGURATION AYZORA AI
# =========================================================

APP_NAME = "AYZORA AI"

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash").strip()

FALLBACK_MODELS = [
    model.strip()
    for model in os.getenv(
        "GEMINI_FALLBACK_MODELS",
        "gemini-3.6-flash,gemini-3.5-flash-lite"
    ).split(",")
    if model.strip()
]

MODEL_CANDIDATES = list(dict.fromkeys([GEMINI_MODEL] + FALLBACK_MODELS))

SUPABASE_URL = os.getenv("SUPABASE_URL", "").rstrip("/")
SUPABASE_KEY = (
    os.getenv("SUPABASE_PUBLISHABLE_KEY", "").strip()
    or os.getenv("SUPABASE_ANON_KEY", "").strip()
)

MAX_FILE_SIZE = 15 * 1024 * 1024
MAX_TEXT_CHARS = 30000
REST_TIMEOUT = 15
GEMINI_RETRIES = 2

client = None

if GEMINI_API_KEY:
    try:
        client = genai.Client(api_key=GEMINI_API_KEY)
    except Exception as exc:
        print("Erreur d'initialisation Gemini :", repr(exc))


# =========================================================
# 2. INSTRUCTIONS DE L'ASSISTANT
# =========================================================

SYSTEM_INSTRUCTION = """
Tu es AYZORA AI, un assistant intelligent polyvalent.

OBJECTIFS :
- Répondre clairement, précisément et utilement.
- Comprendre les besoins de l'utilisateur.
- Aider à analyser des documents et des images.
- Aider à rédiger, apprendre, programmer et résoudre des problèmes.
- Proposer des solutions concrètes et structurées.
- Répondre en français lorsque l'utilisateur écrit en français.

CONTEXTE PROFESSIONNEL :
L'utilisateur peut travailler sur Arrow Home Group et Forza Luxury.
Ne confonds jamais ces deux entreprises.
N'invente pas leurs prix, coordonnées, services ou informations internes.
Utilise les informations disponibles dans les documents fournis,
dans la mémoire persistante et dans la base de connaissances.

MÉMOIRE ET CONFIDENTIALITÉ :
- Utilise uniquement les souvenirs du compte actuellement connecté.
- Ne révèle jamais les souvenirs d'un autre utilisateur.
- Ne prétends jamais avoir mémorisé une information si sa sauvegarde a échoué.
- Si une information manque, demande une précision.

RÈGLES :
- Si une information manque, explique-le honnêtement.
- Ne prétends pas avoir effectué une action que tu n'as pas réalisée.
- Pour le code informatique, fournis des exemples cohérents.
- Pour les documents et les images, analyse réellement leur contenu.
- Sois professionnel, pédagogique et orienté vers les résultats.
"""


# =========================================================
# 3. CONNEXION À SUPABASE
# =========================================================

def supabase_headers(access_token=None, prefer=None):
    if not SUPABASE_KEY:
        return {}

    headers = {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {access_token or SUPABASE_KEY}",
        "Content-Type": "application/json",
    }

    if prefer:
        headers["Prefer"] = prefer

    return headers


def supabase_request(
    method,
    table,
    access_token,
    params=None,
    payload=None,
    prefer=None,
):
    if not SUPABASE_URL or not SUPABASE_KEY:
        print("Supabase non configuré.")
        return None

    try:
        response = requests.request(
            method=method,
            url=f"{SUPABASE_URL}/rest/v1/{table}",
            headers=supabase_headers(access_token, prefer),
            params=params,
            json=payload,
            timeout=REST_TIMEOUT,
        )

        if not response.ok:
            print(
                f"Erreur Supabase {table}:",
                response.status_code,
                response.text[:500],
            )
            return None

        if response.status_code == 204 or not response.content:
            return []

        return response.json()

    except Exception as exc:
        print(f"Erreur Supabase {table} :", repr(exc))
        return None


def get_authenticated_user(access_token):
    if not SUPABASE_URL or not SUPABASE_KEY or not access_token:
        return None

    try:
        response = requests.get(
            f"{SUPABASE_URL}/auth/v1/user",
            headers=supabase_headers(access_token),
            timeout=REST_TIMEOUT,
        )

        if response.ok:
            return response.json()

        print(
            "Vérification utilisateur échouée :",
            response.status_code,
            response.text[:300],
        )

    except Exception as exc:
        print("Erreur vérification utilisateur :", repr(exc))

    return None


# =========================================================
# 4. INSCRIPTION ET CONNEXION
# =========================================================

def auth_request(endpoint, payload):
    if not SUPABASE_URL or not SUPABASE_KEY:
        return None, (
            "Supabase n'est pas configuré. Vérifie les variables "
            "SUPABASE_URL et SUPABASE_PUBLISHABLE_KEY sur Render."
        )

    try:
        response = requests.post(
            f"{SUPABASE_URL}/auth/v1/{endpoint}",
            headers=supabase_headers(),
            json=payload,
            timeout=REST_TIMEOUT,
        )

        data = response.json() if response.content else {}

        if not response.ok:
            error = (
                data.get("msg")
                or data.get("message")
                or data.get("error_description")
                or "Erreur d'authentification."
            )
            return None, error

        return data, ""

    except Exception as exc:
        print("Erreur Supabase Auth :", repr(exc))
        return None, "Impossible de joindre Supabase. Réessaie."


def sign_up(email, password):
    email = (email or "").strip().lower()

    if not email or not password:
        return None, "Saisis une adresse e-mail et un mot de passe."

    if len(password) < 12:
        return None, "Le mot de passe doit contenir au moins 12 caractères."

    data, error = auth_request(
        "signup",
        {
            "email": email,
            "password": password,
        },
    )

    if error:
        return None, error

    if not data.get("access_token"):
        return None, (
            "Inscription envoyée. Vérifie ta boîte e-mail pour confirmer "
            "ton adresse, puis connecte-toi."
        )

    return {
        "access_token": data["access_token"],
        "user": data.get("user", {}),
        "conversation_id": str(uuid.uuid4()),
    }, "Compte créé et connecté."


def sign_in(email, password):
    email = (email or "").strip().lower()

    if not email or not password:
        return None, "Saisis une adresse e-mail et un mot de passe."

    data, error = auth_request(
        "token?grant_type=password",
        {
            "email": email,
            "password": password,
        },
    )

    if error:
        print("Échec connexion :", error)
        return None, (
            "Connexion impossible. Vérifie tes identifiants et "
            "la confirmation de ton adresse e-mail."
        )

    if not data.get("access_token") or not data.get("user"):
        return None, "Supabase n'a pas renvoyé de session valide."

    return {
        "access_token": data["access_token"],
        "user": data["user"],
        "conversation_id": str(uuid.uuid4()),
    }, "Connexion réussie."


# =========================================================
# 5. BASE DE CONNAISSANCES
# =========================================================

def load_knowledge(question, access_token):
    records = supabase_request(
        "GET",
        "ayzora_knowledge",
        access_token,
        params={
            "select": "title,content",
            "verified": "eq.true",
            "limit": "8",
        },
    )

    if not isinstance(records, list):
        return ""

    sections = []

    for record in records:
        title = str(record.get("title") or "").strip()
        content = str(record.get("content") or "").strip()

        if content:
            sections.append(
                f"Titre : {title}\n"
                f"Contenu : {content[:4000]}"
            )

    return "\n\n".join(sections)[:MAX_TEXT_CHARS]


# =========================================================
# MÉMOIRE PERSISTANTE AYZORA AI — VERSION CORRIGÉE
# =========================================================

def load_memories(access_token):
    """Charge uniquement les souvenirs accessibles à l'utilisateur connecté."""
    if not access_token:
        return []

    memories = supabase_request(
        "GET",
        "ayzora_memories",
        access_token,
        params={
            "select": "id,category,title,content,updated_at",
            "order": "updated_at.desc",
            "limit": "50",
        },
    )

    if memories is None:
        print("ERREUR : impossible de lire la mémoire Supabase.")
        return []

    return memories if isinstance(memories, list) else []


def save_memory(access_token, category, title, content):
    """Enregistre un souvenir et vérifie qu'il est réellement sauvegardé."""
    if not access_token:
        return False

    title = str(title or "").strip()
    content = str(content or "").strip()
    category = str(category or "general").strip()

    if not title or not content:
        return False

    user = get_authenticated_user(access_token)

    if not user or not user.get("id"):
        print("ERREUR : utilisateur non authentifié.")
        return False

    payload = {
        "user_id": user["id"],
        "category": category[:80],
        "title": title[:200],
        "content": content[:3000],
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }

    result = supabase_request(
        "POST",
        "ayzora_memories",
        access_token,
        payload=payload,
        prefer="return=representation",
    )

    if result is None:
        print("ÉCHEC : Supabase n'a pas confirmé la sauvegarde.")
        return False

    # Vérifie que le souvenir existe après l'écriture.
    memories = load_memories(access_token)

    for memory in memories:
        if (
            memory.get("title") == title
            and memory.get("content") == content
            and memory.get("user_id", user["id"]) == user["id"]
        ):
            print(f"MÉMOIRE CONFIRMÉE : {title}")
            return True

    # Certaines configurations REST ne renvoient pas la ligne créée.
    # On ne prétend donc pas que la vérification a réussi.
    print(
        "AVERTISSEMENT : écriture envoyée, "
        "mais souvenir non retrouvé dans la vérification."
    )
    return False


def handle_explicit_memory_request(access_token, user_text):
    """
    Reconnaît les demandes explicites comme :
    'Mémorise que mon projet s'appelle BAOBAB-729.'
    """
    import re

    text = (user_text or "").strip()

    if not text:
        return None

    patterns = [
        r"^(?:s'il te plaît[, ]*)?(?:mémorise|memorise|mémoriser|"
        r"retiens|retenir|souviens[- ]toi|enregistre|"
        r"sauvegarde|garde en mémoire|garde en memoire)"
        r"(?:\s+que|\s*:)?\s+(.+)$",
    ]

    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)

        if match:
            information = match.group(1).strip()

            if not information:
                return None

            # Évite d'enregistrer les mots de passe ou les clés.
            sensitive_terms = (
                "mot de passe",
                "password",
                "clé api",
                "cle api",
                "api key",
                "secret key",
                "token d'accès",
            )

            if any(term in information.lower() for term in sensitive_terms):
                return (
                    "Je ne peux pas enregistrer de mot de passe, "
                    "de clé API ou de secret."
                )

            success = save_memory(
                access_token=access_token,
                category="information_utilisateur",
                title="Information à retenir",
                content=information,
            )

            if success:
                return (
                    "✅ **Information enregistrée et vérifiée dans "
                    "la mémoire persistante.**\n\n"
                    f"Je retiendrai : {information}"
                )

            return (
                "⚠️ Je n'ai pas pu confirmer l'enregistrement de cette "
                "information. Vérifie la configuration Supabase et "
                "les logs Render."
            )

    return None


def summarize_and_store_memories(access_token, user_text, assistant_text):
    """Extrait des informations durables sans interrompre le chat."""
    if not client or not user_text.strip():
        return

    prompt = f"""
Extrais au maximum trois informations durables explicitement données
par l'utilisateur, utiles aux futures conversations.

N'invente rien. N'enregistre aucun secret ni mot de passe.
Si aucune information durable n'est présente, renvoie [].

Réponds uniquement avec une liste JSON dont chaque élément contient :
category, title et content.

Message utilisateur :
{user_text[:2500]}

Réponse de l'assistant :
{assistant_text[:1000]}
"""

    try:
        result = generate_with_fallback(
            prompt,
            temperature=0.1,
            max_output_tokens=800,
        )

        raw = (result.text or "").strip()

        if raw.startswith("```"):
            raw = raw.replace("```json", "", 1).replace("```", "").strip()

        start = raw.find("[")
        end = raw.rfind("]")

        if start < 0 or end < start:
            print("Aucun JSON de mémoire valide retourné par Gemini.")
            return

        extracted = json.loads(raw[start:end + 1])

        if not isinstance(extracted, list):
            return

        for item in extracted[:3]:
            if not isinstance(item, dict):
                continue

            title = str(item.get("title") or "").strip()
            content = str(item.get("content") or "").strip()
            category = str(item.get("category") or "general").strip()

            if title and content:
                save_memory(
                    access_token,
                    category,
                    title,
                    content,
                )

    except Exception as exc:
        print("Erreur extraction mémoire :", repr(exc))

# =========================================================
# 7. GEMINI AVEC NOUVELLES TENTATIVES ET MODÈLES DE SECOURS
# =========================================================

def generate_with_fallback(
    contents,
    temperature=0.4,
    max_output_tokens=4096,
    model_candidates=None,
):
    if not client:
        raise RuntimeError("La clé GEMINI_API_KEY n'est pas configurée.")

    candidates = model_candidates or MODEL_CANDIDATES
    errors = []

    for model_name in list(dict.fromkeys(candidates)):
        for attempt in range(GEMINI_RETRIES):
            try:
                print(
                    f"Appel Gemini : modèle={model_name}, "
                    f"tentative={attempt + 1}"
                )

                return client.models.generate_content(
                    model=model_name,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        temperature=temperature,
                        max_output_tokens=max_output_tokens,
                    ),
                )

            except Exception as exc:
                error_message = str(exc)
                errors.append(
                    f"{model_name}: {type(exc).__name__}: "
                    f"{error_message[:180]}"
                )

                text = error_message.lower()

                retryable = any(
                    code in text
                    for code in (
                        "503",
                        "unavailable",
                        "429",
                        "resource_exhausted",
                        "500",
                        "internal",
                        "502",
                        "504",
                        "deadline",
                        "timed out",
                        "timeout",
                    )
                )

                print(
                    f"Échec Gemini {model_name}, "
                    f"tentative {attempt + 1} :",
                    repr(exc),
                )

                if not retryable:
                    if any(
                        term in text
                        for term in (
                            "not found",
                            "unsupported",
                            "model",
                        )
                    ):
                        break
                    raise

                if attempt < GEMINI_RETRIES - 1:
                    time.sleep(1.5 * (2 ** attempt))
                else:
                    break

    raise RuntimeError(
        "Tous les modèles Gemini ont échoué. "
        + " | ".join(errors[-6:])
    )


def summarize_and_store_memories(access_token, user_text, assistant_text):
    """
    Extrait au maximum trois informations durables.
    Une panne de mémoire ne doit pas interrompre le chat.
    """

    if not client or not user_text.strip():
        return

    prompt = f"""
Tu es un extracteur de mémoire pour un assistant personnel.

Extrais au maximum trois informations durables explicitement données
par l'utilisateur et utiles pour de futures conversations.

Règles :
- N'infère aucun fait.
- N'enregistre aucun mot de passe, clé API ou secret.
- N'enregistre pas de données sensibles.
- N'enregistre pas une question ponctuelle comme un souvenir.
- Si rien n'est durable, renvoie [].
- Chaque souvenir doit être compréhensible sans le contexte.
- Réponds uniquement en JSON sous forme d'une liste.
- Les clés obligatoires sont category, title et content.

Message utilisateur :
{user_text[:2500]}

Réponse de l'assistant :
{assistant_text[:1000]}
"""

    try:
        result = generate_with_fallback(
            prompt,
            temperature=0.1,
            max_output_tokens=800,
        )

        raw = (result.text or "").strip()

        if raw.startswith("```"):
            raw = raw.replace("```json", "", 1).replace("```", "").strip()

        start = raw.find("[")
        end = raw.rfind("]")

        if start == -1 or end == -1:
            return

        memories = json.loads(raw[start:end + 1])

        if not isinstance(memories, list):
            return

        for item in memories[:3]:
            if not isinstance(item, dict):
                continue

            title = str(item.get("title") or "").strip()
            content = str(item.get("content") or "").strip()
            category = str(item.get("category") or "general").strip()

            if title and content:
                save_memory(
                    access_token,
                    category,
                    title,
                    content,
                )

    except Exception as exc:
        print("Extraction de mémoire ignorée :", repr(exc))


def refresh_memory(session):
    if not session or not session.get("access_token"):
        return "Connecte-toi pour consulter ta mémoire."

    memories = load_memories(session["access_token"])

    if not memories:
        return (
            "Aucun souvenir enregistré pour le moment.\n\n"
            "Tu peux demander à AYZORA de retenir une préférence ou "
            "une information durable dans une conversation."
        )

    sections = []

    for memory in memories:
        sections.append(
            f"**{memory.get('title', 'Souvenir')}**\n"
            f"_Catégorie : {memory.get('category', 'general')}_\n\n"
            f"{memory.get('content', '')}"
        )

    return "\n\n---\n\n".join(sections)


# =========================================================
# 8. HISTORIQUE DES CONVERSATIONS
# =========================================================

def save_conversation(access_token, conversation_id, messages, title=None):
    user = get_authenticated_user(access_token)

    if not user or not user.get("id"):
        print("Conversation non sauvegardée : utilisateur non authentifié.")
        return False

    payload = {
        "id": conversation_id,
        "user_id": user["id"],
        "title": (title or "Conversation AYZORA")[:200],
        "messages": messages[-80:],
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }

    result = supabase_request(
        "POST",
        "ayzora_conversations",
        access_token,
        params={"on_conflict": "id"},
        payload=payload,
        prefer="resolution=merge-duplicates,return=minimal",
    )

    return result is not None


def get_conversation_choices(session):
    if not session or not session.get("access_token"):
        return []

    records = supabase_request(
        "GET",
        "ayzora_conversations",
        session["access_token"],
        params={
            "select": "id,title,updated_at",
            "order": "updated_at.desc",
            "limit": "50",
        },
    )

    if not isinstance(records, list):
        return []

    return [
        (
            f"{record.get('title') or 'Conversation'} — "
            f"{str(record.get('updated_at') or '')[:16]}",
            record["id"],
        )
        for record in records
        if record.get("id")
    ]


def refresh_conversation_choices(session):
    return gr.update(
        choices=get_conversation_choices(session),
        value=None,
    )


def load_conversation(session, conversation_id):
    if not session or not session.get("access_token"):
        return [], None, "Connecte-toi pour ouvrir une conversation."

    if not conversation_id:
        return [], None, "Choisis une conversation dans la liste."

    records = supabase_request(
        "GET",
        "ayzora_conversations",
        session["access_token"],
        params={
            "select": "id,title,messages",
            "id": f"eq.{conversation_id}",
            "limit": "1",
        },
    )

    if not isinstance(records, list) or not records:
        return [], None, "Conversation introuvable."

    record = records[0]
    messages = record.get("messages") or []

    if not isinstance(messages, list):
        messages = []

    cleaned = []

    for item in messages:
        if (
            isinstance(item, dict)
            and item.get("role") in ("user", "assistant")
        ):
            content = item.get("content", "")

            if isinstance(content, str):
                cleaned.append(
                    {
                        "role": item["role"],
                        "content": content,
                    }
                )

    return (
        cleaned,
        record.get("id"),
        f"Conversation ouverte : **{record.get('title') or 'Conversation'}**",
    )


def new_conversation(session):
    new_id = str(uuid.uuid4())

    if session:
        session["conversation_id"] = new_id

    return [], new_id, "Nouvelle conversation prête."


def login_action(email, password, mode):
    if mode == "Créer un compte":
        session, message = sign_up(email, password)
    else:
        session, message = sign_in(email, password)

    if not session:
        return (
            None,
            f"**{message}**",
            gr.update(visible=True),
            gr.update(visible=False),
            [],
            None,
            gr.update(choices=[], value=None),
        )

    if not session.get("user", {}).get("id"):
        return (
            None,
            "**Impossible de vérifier le compte utilisateur.**",
            gr.update(visible=True),
            gr.update(visible=False),
            [],
            None,
            gr.update(choices=[], value=None),
        )

    choices = get_conversation_choices(session)

    return (
        session,
        f"**{message}**",
        gr.update(visible=False),
        gr.update(visible=True),
        [],
        session["conversation_id"],
        gr.update(choices=choices, value=None),
    )


def logout_action():
    return (
        None,
        "**Vous êtes déconnecté.**",
        gr.update(visible=True),
        gr.update(visible=False),
        [],
        None,
        gr.update(choices=[], value=None),
        "",
    )


# =========================================================
# 9. EXTRACTION DU TEXTE DES DOCUMENTS
# =========================================================

def extract_document_text(file_path):
    path = Path(file_path)
    extension = path.suffix.lower()

    if not path.exists():
        return "[Fichier introuvable]"

    if path.stat().st_size > MAX_FILE_SIZE:
        return f"[Le fichier {path.name} dépasse la limite de 15 Mo.]"

    try:
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))

            return "\n".join(
                page.extract_text() or ""
                for page in reader.pages[:40]
            )[:MAX_TEXT_CHARS]

        if extension == ".docx":
            from docx import Document

            document = Document(str(path))

            return "\n".join(
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            )[:MAX_TEXT_CHARS]

        if extension in {".txt", ".md", ".log"}:
            return path.read_text(
                encoding="utf-8",
                errors="replace",
            )[:MAX_TEXT_CHARS]

        if extension == ".json":
            with path.open(
                "r",
                encoding="utf-8",
                errors="replace",
            ) as file:
                data = json.load(file)

            return json.dumps(
                data,
                ensure_ascii=False,
                indent=2,
            )[:MAX_TEXT_CHARS]

        if extension == ".csv":
            with path.open(
                "r",
                encoding="utf-8-sig",
                errors="replace",
                newline="",
            ) as file:
                rows = list(csv.reader(file))

            return "\n".join(
                " | ".join(row)
                for row in rows[:300]
            )[:MAX_TEXT_CHARS]

        if extension in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(
                str(path),
                read_only=True,
                data_only=True,
            )

            sections = []
            total = 0

            try:
                for sheet in workbook.worksheets:
                    sections.append(f"\nFeuille : {sheet.title}")

                    for row in sheet.iter_rows(
                        min_row=1,
                        max_row=150,
                        values_only=True,
                    ):
                        line = " | ".join(
                            "" if value is None else str(value)
                            for value in row
                        )

                        sections.append(line)
                        total += len(line)

                        if total >= MAX_TEXT_CHARS:
                            break

                    if total >= MAX_TEXT_CHARS:
                        break

            finally:
                workbook.close()

            return "\n".join(sections)[:MAX_TEXT_CHARS]

        return (
            f"[Le format {extension} n'est pas pris en charge "
            "pour l'extraction de texte.]"
        )

    except Exception as exc:
        print(
            f"Erreur de lecture du fichier {path.name} :",
            repr(exc),
        )
        return f"[Impossible de lire le fichier {path.name}.]"


# =========================================================
# 10. PRÉPARATION DES FICHIERS POUR GEMINI
# =========================================================

def prepare_file_parts(file_paths):
    parts = []

    for file_path in file_paths or []:
        if not file_path:
            continue

        path = Path(file_path)

        if not path.exists():
            parts.append(
                types.Part.from_text(
                    text=f"Fichier introuvable : {path.name}"
                )
            )
            continue

        if path.stat().st_size > MAX_FILE_SIZE:
            parts.append(
                types.Part.from_text(
                    text=f"Fichier trop volumineux : {path.name}"
                )
            )
            continue

        if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
            mime_type = (
                mimetypes.guess_type(str(path))[0] or "image/jpeg"
            )

            parts.append(
                types.Part.from_text(
                    text=f"Image fournie : {path.name}"
                )
            )

            parts.append(
                types.Part.from_bytes(
                    data=path.read_bytes(),
                    mime_type=mime_type,
                )
            )

        else:
            extracted_text = extract_document_text(path)

            parts.append(
                types.Part.from_text(
                    text=(
                        f"Document : {path.name}\n"
                        f"Contenu extrait :\n{extracted_text}"
                    )
                )
            )

    return parts


# =========================================================
# 11. FONCTION PRINCIPALE DU CHAT
# =========================================================

def chat(message, history, session, conversation_id):
    if not session or not session.get("access_token"):
        return "Connecte-toi pour utiliser AYZORA AI."

    if not client:
        return "La clé GEMINI_API_KEY n'est pas configurée sur Render."

    if isinstance(message, dict):
        question = str(message.get("text") or "").strip()
        file_paths = message.get("files") or []
    else:
        question = str(message or "").strip()
        file_paths = []

    if isinstance(file_paths, str):
        file_paths = [file_paths]

    if not question and not file_paths:
        return "Écris une question ou joins un fichier."

    token = session["access_token"]

    try:
        user = get_authenticated_user(token)

        if not user:
            return (
                "Ta session a expiré ou n'est plus valide. "
                "Déconnecte-toi, puis reconnecte-toi."
            )

        memories = load_memories(token)
        memory_text = retrieve_memories_for_prompt(memories)
        knowledge = load_knowledge(question, token)

        contents = []

        for item in (history or [])[-12:]:
            if not isinstance(item, dict):
                continue

            role = item.get("role")
            content = item.get("content")

            if (
                role in {"user", "assistant"}
                and isinstance(content, str)
                and content.strip()
            ):
                contents.append(
                    types.Content(
                        role="user" if role == "user" else "model",
                        parts=[
                            types.Part.from_text(text=content)
                        ],
                    )
                )

        prompt_parts = [
            types.Part.from_text(text=SYSTEM_INSTRUCTION)
        ]

        if memory_text:
            prompt_parts.append(
                types.Part.from_text(
                    text=(
                        "MÉMOIRE PERSISTANTE DE CET UTILISATEUR :\n"
                        f"{memory_text}\n\n"
                        "Utilise ces souvenirs uniquement s'ils sont "
                        "pertinents pour la demande actuelle."
                    )
                )
            )

        if knowledge:
            prompt_parts.append(
                types.Part.from_text(
                    text=(
                        "BASE DE CONNAISSANCES VÉRIFIÉE :\n"
                        f"{knowledge}\n\n"
                        "Utilise ces informations lorsqu'elles sont pertinentes."
                    )
                )
            )

        prompt_parts.append(
            types.Part.from_text(
                text=(
                    "MESSAGE DE L'UTILISATEUR :\n"
                    f"{question or 'Analyse les fichiers joints.'}"
                )
            )
        )

        prompt_parts.extend(prepare_file_parts(file_paths))

        contents.append(
            types.Content(
                role="user",
                parts=prompt_parts,
            )
        )

        response = generate_with_fallback(contents)
        answer = (response.text or "").strip()

        if not answer:
            answer = (
                "Je n'ai pas pu produire de réponse textuelle. "
                "Essaie de reformuler ta question."
            )

        current_history = list(history or [])

        current_history.extend([
            {
                "role": "user",
                "content": question or "[Fichier joint]",
            },
            {
                "role": "assistant",
                "content": answer,
            },
        ])

        if not conversation_id:
            conversation_id = (
                session.get("conversation_id") or str(uuid.uuid4())
            )

        session["conversation_id"] = conversation_id

        saved = save_conversation(
            token,
            conversation_id,
            current_history,
            title=(question or "Analyse de fichier")[:100],
        )

        if not saved:
            print("Attention : la conversation n'a pas été sauvegardée.")

        if question:
            summarize_and_store_memories(
                token,
                question,
                answer,
            )

        return answer

    except Exception as exc:
        print("ERREUR AYZORA AI :", repr(exc))
        traceback.print_exc()

        error_text = str(exc).lower()

        if "503" in error_text or "unavailable" in error_text:
            return (
                "Gemini est temporairement indisponible, même après les "
                "tentatives et modèles de secours. Attends un peu puis réessaie."
            )

        if "429" in error_text or "resource_exhausted" in error_text:
            return (
                "La limite de requêtes Gemini a été atteinte. "
                "Attends avant de réessayer."
            )

        return (
            f"Une erreur est survenue ({type(exc).__name__}). "
            "Consulte les logs Render pour le détail."
        )


# =========================================================
# 12. INTERFACE UTILISATEUR
# =========================================================

DESCRIPTION = """
### Bienvenue sur AYZORA AI

Ton assistant intelligent pour :
- Répondre à tes questions et t'aider à apprendre.
- Rédiger des contenus et préparer des stratégies.
- Analyser des images et des documents.
- T'accompagner dans tes projets professionnels et techniques.

Connecte-toi pour accéder à tes conversations et à ta mémoire personnelle.
"""


with gr.Blocks(
    theme=gr.themes.Soft(),
    title="AYZORA AI",
) as demo:

    session_state = gr.State(None)
    conversation_state = gr.State(None)

    gr.Markdown(
        "# 🟠 AYZORA AI\n"
        "### Ton assistant intelligent"
    )

    gr.Markdown(DESCRIPTION)

    with gr.Column(visible=True) as auth_panel:
        gr.Markdown("## Connexion / inscription")

        mode = gr.Radio(
            ["Se connecter", "Créer un compte"],
            value="Se connecter",
            label="Action",
        )

        email = gr.Textbox(
            label="Adresse e-mail",
            placeholder="vous@exemple.com",
        )

        password = gr.Textbox(
            label="Mot de passe (12 caractères minimum)",
            type="password",
        )

        auth_button = gr.Button(
            "Continuer",
            variant="primary",
        )

        auth_status = gr.Markdown(
            "Les nouveaux comptes peuvent nécessiter une confirmation par e-mail."
        )

    with gr.Column(visible=False) as app_panel:
        login_status = gr.Markdown()

        with gr.Row():
            new_chat_button = gr.Button(
                "➕ Nouvelle conversation",
                variant="primary",
            )

            memory_button = gr.Button("🧠 Voir ma mémoire")
            logout_button = gr.Button("Se déconnecter")

        with gr.Row():
            conversation_picker = gr.Dropdown(
                label="Conversations précédentes",
                choices=[],
                value=None,
                interactive=True,
                scale=4,
            )

            refresh_history_button = gr.Button("↻ Actualiser")
            open_history_button = gr.Button("Ouvrir")

        history_status = gr.Markdown()
        memory_view = gr.Markdown()

        chatbot = gr.ChatInterface(
            fn=chat,
            type="messages",
            multimodal=True,
            chatbot=gr.Chatbot(
                type="messages",
                height=500,
                placeholder="Pose ta question à AYZORA AI...",
            ),
            textbox=gr.MultimodalTextbox(
                placeholder="Écris ta question ou joins un fichier...",
                file_count="multiple",
                file_types=[
                    ".pdf",
                    ".docx",
                    ".txt",
                    ".md",
                    ".log",
                    ".json",
                    ".csv",
                    ".xlsx",
                    ".xlsm",
                    ".png",
                    ".jpg",
                    ".jpeg",
                    ".webp",
                ],
                sources=["upload"],
            ),
            additional_inputs=[
                session_state,
                conversation_state,
            ],
            examples=[
                [
                    "Présente-toi et explique-moi tes capacités.",
                    None,
                    None,
                ],
                [
                    "Aide-moi à préparer une stratégie marketing digital.",
                    None,
                    None,
                ],
                [
                    "Explique-moi simplement ce qu'est un agent IA.",
                    None,
                    None,
                ],
                [
                    "Je souhaite que tu retiennes que mon objectif est de développer AYZORA AI.",
                    None,
                    None,
                ],
            ],
            save_history=False,
        )

    auth_button.click(
        fn=login_action,
        inputs=[email, password, mode],
        outputs=[
            session_state,
            auth_status,
            auth_panel,
            app_panel,
            chatbot.chatbot,
            conversation_state,
            conversation_picker,
        ],
    )

    logout_button.click(
        fn=logout_action,
        inputs=[],
        outputs=[
            session_state,
            auth_status,
            auth_panel,
            app_panel,
            chatbot.chatbot,
            conversation_state,
            conversation_picker,
            memory_view,
        ],
    )

    memory_button.click(
        fn=refresh_memory,
        inputs=[session_state],
        outputs=[memory_view],
    )

    refresh_history_button.click(
        fn=refresh_conversation_choices,
        inputs=[session_state],
        outputs=[conversation_picker],
    )

    open_history_button.click(
        fn=load_conversation,
        inputs=[
            session_state,
            conversation_picker,
        ],
        outputs=[
            chatbot.chatbot,
            conversation_state,
            history_status,
        ],
    )

    new_chat_button.click(
        fn=new_conversation,
        inputs=[session_state],
        outputs=[
            chatbot.chatbot,
            conversation_state,
            history_status,
        ],
    )

    gr.Markdown(
        "---\n"
        "AYZORA AI · Assistant intelligent"
    )


# =========================================================
# 13. DÉMARRAGE
# =========================================================

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "7860"))

    demo.launch(
        server_name="0.0.0.0",
        server_port=port,
    )
